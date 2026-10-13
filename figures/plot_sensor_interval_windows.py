"""单传感器两个月滑窗的连续光滑概率轮廓，只输出一张Matplotlib SVG。
默认窗口2mo、步长2mo，按自然月计算；也支持30d、15d等固定时长。
默认保存到figures/interval_scattering；图中保留编码与含义标题、曲线、坐标轴和图例。
修改DEFAULTS或用命令行参数覆盖。高度是窗口内出现概率，不是每秒密度。
依赖numpy、pandas、matplotlib、scipy；读取选取点表时另需openpyxl。
"""
from __future__ import annotations
import argparse
import hashlib
import json
import re
from pathlib import Path
import sys
import tempfile
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import PercentFormatter
import scipy
from scipy.interpolate import PchipInterpolator

ROOT = Path(__file__).resolve().parents[1]
SECOND = 1_000_000_000
DEFAULTS = dict(
    unit="jqrd_gt1", sensor="JQRD_11EKG60FF001OUT", name=None,
    start=None, end=None, window="2mo", step="2mo",
    state_mode="all", power_threshold=1.0, state_max_gap=5.0,
    min_running_fraction=0.5, min_samples=50,
    bin_scale="log", bins=80, min_interval=0.001, max_interval=604800.0,
    chunk_size=350000, slices=0, smooth_points=2400,
    output_dir=ROOT / "figures" / "interval_scattering",
)

def time_ns(value):
    t = pd.Timestamp(value)
    if t.tzinfo is not None:
        t = t.tz_convert("Asia/Shanghai").tz_localize(None)
    return t.value


def timestamps(column):
    t = pd.to_datetime(column, format="mixed", errors="raise")
    if t.isna().any():
        raise ValueError("存在空时间戳；不能静默猜测顺序。")
    if isinstance(t.dtype, pd.DatetimeTZDtype):
        t = t.dt.tz_convert("Asia/Shanghai").dt.tz_localize(None)
    return t.astype("int64").to_numpy()


def metadata(path):
    stat = path.stat()
    return dict(path=str(path.resolve()), bytes=stat.st_size, mtime_ns=stat.st_mtime_ns)


def chunks(path, size):
    """保持文件物理行顺序；不排序，不插值。"""
    previous = None
    for c in pd.read_csv(path, usecols=["timestamp", "value"], chunksize=size):
        t = timestamps(c["timestamp"])
        v = pd.to_numeric(c["value"], errors="coerce").to_numpy(dtype=float)
        if len(t) and (np.any(np.diff(t) < 0) or (previous is not None and t[0] < previous)):
            raise ValueError(f"时间戳倒序：{path}；请先核查来源。")
        if len(t):
            previous = int(t[-1])
            yield t, v


def selected_points(unit):
    from openpyxl import load_workbook
    files = sorted((ROOT / "raw_data" / unit).glob("*选取点表.xlsx"))
    if not files:
        return [], None
    if len(files) != 1:
        raise ValueError("存在多个选取点表；请用 --input、--name 和 --power-csv 明确指定。")
    wb = load_workbook(files[0], read_only=True, data_only=True)
    try:
        points = [tuple(r) for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True) if r[0]]
    finally:
        wb.close()
    return points, files[0]


def power_spans(path, threshold, max_gap, chunk_size):
    """仅连接相邻、有效、均在运行且间隔不超过 max_gap 的出力记录。"""
    starts, ends = [], []
    previous = None
    rows = 0
    for t, v in chunks(path, chunk_size):
        rows += len(t)
        if previous is not None:
            t = np.r_[previous[0], t]
            v = np.r_[previous[1], v]
        previous = (int(t[-1]), float(v[-1]))
        good = np.isfinite(v) & (v > threshold)
        dt = np.diff(t)
        keep = (dt > 0) & (dt <= max_gap * SECOND) & good[:-1] & good[1:]
        a = np.flatnonzero(keep & ~np.r_[False, keep[:-1]])
        b = np.flatnonzero(keep & ~np.r_[keep[1:], False]) + 1
        for left, right in zip(t[a], t[b]):
            if ends and left <= ends[-1]:
                ends[-1] = max(ends[-1], int(right))
            else:
                starts.append(int(left)); ends.append(int(right))
    return np.array(starts, dtype=np.int64), np.array(ends, dtype=np.int64), rows


def span_ids(t, starts, ends):
    if not len(starts):
        return np.full(len(t), -1, dtype=np.int64)
    i = np.searchsorted(starts, t, side="right") - 1
    return np.where((i >= 0) & (t <= ends[np.maximum(i, 0)]), i, -1)


def support_before(q, starts, ends):
    prefix = np.r_[0.0, np.cumsum((ends - starts) / SECOND)]
    i = np.searchsorted(ends, q, side="right")
    out = prefix[i].copy()
    available = i < len(starts)
    out[available] += np.maximum(0, q[available] - starts[i[available]]) / SECOND
    return out


def add_windows(diff, left, categories, origin, window, step, n, window_bounds=None):
    """用差分累加所有包含间隔左端的窗口，不复制重叠窗口样本。"""
    if window_bounds is None:
        lo = np.maximum(0, (left - origin - window) // step + 1)
        hi = np.minimum(n - 1, (left - origin) // step)
    else:
        starts,stops=window_bounds
        lo=np.searchsorted(stops,left,side="right")
        hi=np.searchsorted(starts,left,side="right")-1
    valid = lo <= hi
    np.add.at(diff, (lo[valid], categories[valid]), 1)
    np.add.at(diff, (hi[valid] + 1, categories[valid]), -1)


def scan_intervals(path, start, end, window, step, edges, spans, chunk_size, window_bounds=None):
    n = len(window_bounds[0]) if window_bounds is not None else int((end - start - 1) // step + 1)
    # 类别：下溢、各可见箱、上溢、跨状态边界排除、数据末端右删失。
    nb = len(edges) - 1
    diff = np.zeros((n + 1, nb + 4), dtype=np.int64)
    previous = None
    audit = dict(rows=0, invalid_values=0, duplicate_valid_timestamps=0,
                 positive_intervals=0, intervals_starting_in_range=0,
                 state_excluded_intervals=0, eligible_intervals=0)
    for t, v in chunks(path, chunk_size):
        audit["rows"] += len(t)
        valid = np.isfinite(v)
        audit["invalid_values"] += int((~valid).sum())
        t = t[valid]
        if not len(t):
            continue
        if previous is not None:
            t = np.r_[previous, t]
        left, right = t[:-1], t[1:]
        positive = right > left
        audit["duplicate_valid_timestamps"] += int((~positive).sum())
        left, right = left[positive], right[positive]
        audit["positive_intervals"] += len(left)
        in_range = (left >= start) & (left < end)
        left, right = left[in_range], right[in_range]
        audit["intervals_starting_in_range"] += len(left)
        if spans is not None:
            a = span_ids(left, *spans); b = span_ids(right, *spans)
            good = (a >= 0) & (a == b)
            audit["state_excluded_intervals"] += int((~good).sum())
            add_windows(diff, left[~good], np.full(int((~good).sum()), nb + 2), start, window, step, n, window_bounds)
            left, right = left[good], right[good]
        delta = (right - left) / SECOND
        categories = np.searchsorted(edges, delta, side="right")
        categories[delta == edges[-1]] = nb  # 最后一个可见箱包含其右边界。
        add_windows(diff, left, categories, start, window, step, n, window_bounds)
        audit["eligible_intervals"] += len(left)
        previous = int(t[-1])
    # 只标记数据末端没有下一条有效观测的记录；不把它作为已完成间隔。
    if previous is not None and start <= previous < end:
        supported = spans is None or span_ids(np.array([previous]), *spans)[0] >= 0
        if supported:
            add_windows(diff, np.array([previous]), np.array([nb + 3]), start, window, step, n, window_bounds)
    hist = np.cumsum(diff[:-1], axis=0)
    assert np.all(hist >= 0)
    return hist, audit


def source_extent(path):
    """从首行及末行读范围；输入为本项目 timestamp,value 单行 CSV。"""
    import csv
    first = pd.read_csv(path, nrows=1)
    if first.empty:
        raise ValueError("输入 CSV 没有记录。")
    with path.open("rb") as f:
        f.seek(0, 2); size = f.tell()
        width = min(size, 65536)
        f.seek(size - width)
        lines = f.read().decode("utf-8-sig").splitlines()
    last = next(line for line in reversed(lines) if line.strip())
    record = next(csv.reader([last]))
    index = first.columns.tolist().index("timestamp")
    return time_ns(first.iloc[0]["timestamp"]), time_ns(record[index])


def self_check():
    with tempfile.TemporaryDirectory() as td:
        source = Path(td) / "sensor.csv"
        base = time_ns("2025-01-01")
        t = [0, 1, 1, 2, 3, 4, 10]
        pd.DataFrame({"timestamp": pd.to_datetime(base + np.array(t) * SECOND),
                      "value": [1, 1, 2, 1, np.nan, 1, 1]}).to_csv(source, index=False)
        edges = np.array([.5, 1.5, 3, 5.0])
        hist, audit = scan_intervals(source, base, base + 7 * SECOND, 4 * SECOND, 2 * SECOND, edges, None, 2)
        # 有效相邻间隔：1,1,2,6；最后6秒间隔跨过分析终点，仍按左端归窗。
        np.testing.assert_array_equal(hist[:, :5], [[0, 2, 1, 0, 0], [0, 0, 1, 0, 1], [0, 0, 0, 0, 1], [0, 0, 0, 0, 0]])
        assert audit["duplicate_valid_timestamps"] == 1 and audit["invalid_values"] == 1
        support = (base + np.array([0, 4]) * SECOND, base + np.array([2, 10]) * SECOND)
        filtered, _ = scan_intervals(source, base, base + 11 * SECOND, 4 * SECOND, 2 * SECOND, edges, support, 3)
        assert filtered[0, 5] == 1  # 2->4 跨运行支持边界。
        assert filtered[4, 6] == 1  # 最后一条有效记录右删失。
        q = base + np.array([0, 2, 3, 4, 7, 11]) * SECOND
        np.testing.assert_allclose(support_before(q, *support), [0, 2, 2, 2, 5, 8])
        pd.DataFrame({"timestamp": pd.to_datetime(base + np.array([0, 1, 2, 3, 4, 8, 9]) * SECOND),
                      "value": [2, 2, np.nan, 2, 2, 2, 2]}).to_csv(source, index=False)
        s, e, _ = power_spans(source, 1, 2, 2)
        np.testing.assert_array_equal((s - base) // SECOND, [0, 3, 8])
        np.testing.assert_array_equal((e - base) // SECOND, [1, 4, 9])
        # 分块大小不改变结果。
        s2, e2, _ = power_spans(source, 1, 2, 100)
        np.testing.assert_array_equal(s, s2); np.testing.assert_array_equal(e, e2)
    e=np.array([.01,.1,1.,10.,100.])
    p=np.array([0.,.8,.1,0.])
    for scale in ("log","linear"):
        x,y=smooth_probability(e,p,scale,2400)
        assert np.all(np.isfinite(y)) and np.all(y>=0) and np.all(y<=1)
    assert select_slices(23,0).tolist()==list(range(23))
    print("SELF_CHECK_PASSED: interval boundaries, probabilities, censoring, positive smooth interpolation")




def parse_period(spec):
    match=re.fullmatch(r"([1-9][0-9]*)mo",spec.strip().lower())
    if match:
        return "months",int(match.group(1))
    duration=int(pd.Timedelta(spec).value)
    if duration<=0:
        raise ValueError("窗口或步长须为正；自然月使用3mo等格式。")
    return "duration",duration


def period_end(start,spec):
    kind,amount=parse_period(spec)
    return int((pd.Timestamp(start)+pd.DateOffset(months=amount)).value) if kind=="months" else int(start)+amount


def period_label(spec):
    kind,amount=parse_period(spec)
    return f"{amount}个月" if kind=="months" else spec


def make_windows(start,end,window_spec,step_spec,max_windows):
    step_kind,amount=parse_period(step_spec)
    if step_kind=="duration":
        n=int((end-start-1)//amount+1)
        if n>max_windows:
            raise ValueError("滑窗统计超过上限；请增加step或减少bins。")
        starts=start+np.arange(n,dtype=np.int64)*amount
    else:
        anchor=pd.Timestamp(start)
        values=[]
        k=0
        while True:
            point=int((anchor+pd.DateOffset(months=k*amount)).value)
            if point>=end:break
            if len(values)>=max_windows:
                raise ValueError("滑窗统计超过上限；请增加step或减少bins。")
            values.append(point);k+=1
        starts=np.array(values,dtype=np.int64)
    stops=np.array([min(period_end(int(a),window_spec),end) for a in starts],dtype=np.int64)
    if np.any(stops<=starts) or np.any(np.diff(starts)<=0) or np.any(np.diff(stops)<0):
        raise ValueError("窗口边界必须有序且长度为正。")
    return starts,stops


def select_slices(window_count, requested):
    """默认画每个窗口；requested>0时按时间等间距取指定数量。"""
    if window_count<1 or requested<0:
        raise ValueError("滑窗数须为正，切片数不能为负。")
    if requested==0:
        return np.arange(window_count)
    return np.unique(np.rint(np.linspace(0,window_count-1,min(requested,window_count))).astype(int))


def smooth_probability(edges, probabilities, scale, points):
    """保形PCHIP平滑显示分箱概率；保留各箱中心概率，不作为新增观测或密度。"""
    mid=np.sqrt(edges[:-1]*edges[1:]) if scale=="log" else (edges[:-1]+edges[1:])/2
    knots=np.r_[edges[0],mid,edges[-1]]
    values=np.r_[probabilities[0],probabilities,probabilities[-1]]
    coordinates=np.log10(knots) if scale=="log" else knots
    grid=np.linspace(coordinates[0],coordinates[-1],points)
    interpolator=PchipInterpolator(coordinates,values,extrapolate=False)
    curve=interpolator(grid)
    if np.any(curve < -1e-12) or np.any(curve > 1+1e-12):
        raise ValueError("平滑曲线超出概率范围。")
    np.testing.assert_allclose(interpolator(coordinates),values,atol=1e-12)
    return 10**grid if scale=="log" else grid,np.clip(curve,0,1)


def write_slices(output,starts,stops,edges,probabilities,bin_counts,counts,usable,name,args,provenance):
    plt.rcParams.update({"font.sans-serif":["Microsoft YaHei","SimHei","DejaVu Sans"],
                         "axes.unicode_minus":False,"svg.fonttype":"path"})
    selected=select_slices(len(starts),args.slices)
    height=max(8.5,2.5+.27*len(selected))
    fig,ax=plt.subplots(figsize=(16,height))
    fig.subplots_adjust(left=.075,right=.715,bottom=.11,top=.90)
    colours=plt.get_cmap("viridis")(np.linspace(.05,.95,len(selected)))
    windows=[]
    for number,index in enumerate(selected):
        left=pd.Timestamp(starts[index]);right=pd.Timestamp(stops[index])
        label=left.strftime("%Y-%m-%d")+" 至 "+right.strftime("%Y-%m-%d")
        windows.append(dict(index=int(index),start=str(left),end=str(right),
                            intervals=int(counts[index]),usable=bool(usable[index])))
        if not usable[index]:
            continue
        x,y=smooth_probability(edges,probabilities[index],args.bin_scale,args.smooth_points)
        ax.plot(x,y,color=colours[number],lw=1.9,ls="-",alpha=.88,label=label)
    if args.bin_scale=="log":
        ax.set_xscale("log")
    ax.set_xlim(edges[0],edges[-1]);ax.set_ylim(0,1.03)
    ax.yaxis.set_major_formatter(PercentFormatter(xmax=1))
    ax.set_xlabel("相邻有效测量的时间间隔 / 秒",fontsize=12,labelpad=9)
    ax.set_ylabel("窗口内出现概率",fontsize=12,labelpad=9)
    ax.set_title(provenance.get("sensor",args.sensor)+"  "+name,fontsize=14,pad=16)
    ax.grid(which="major",color="#cbd5e1",alpha=.55,lw=.65)
    ax.spines[["top","right"]].set_visible(False)
    if usable[selected].any():
        ax.legend(loc="upper left",bbox_to_anchor=(1.02,1),
                  frameon=False,fontsize=8.7,handlelength=2.7,labelspacing=.72)
    provenance["windows"]=[dict(start=str(pd.Timestamp(a)),end=str(pd.Timestamp(b)),intervals=int(n),usable=bool(ok))
                           for a,b,n,ok in zip(starts,stops,counts,usable)]
    provenance["selected_slices"]=windows
    provenance["smoothing"]="Shape-preserving PCHIP of bin probabilities in log-interval coordinates for log scale, linear coordinates otherwise. C1 continuous, nonnegative; centre probabilities unchanged. Visual interpolation only, not a continuous probability density or fitted theoretical distribution."
    provenance["all_window_probabilities"]=np.where(np.isfinite(probabilities),probabilities,None).tolist()
    provenance["all_window_bin_counts"]=bin_counts.tolist()
    description=json.dumps(provenance,ensure_ascii=False,allow_nan=False)
    fig.savefig(output,format="svg",metadata={"Title":provenance.get("sensor",args.sensor)+"  "+name,
                                            "Description":description,"Creator":"Matplotlib"})
    plt.close(fig)


def parser():
    p=argparse.ArgumentParser(description=__doc__,formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--input",type=Path,help="自定义timestamp,value CSV")
    p.add_argument("--unit",default=DEFAULTS["unit"])
    p.add_argument("--sensor",default=DEFAULTS["sensor"])
    p.add_argument("--name",default=DEFAULTS["name"])
    p.add_argument("--list-sensors",action="store_true")
    for key in ("start","end","window","step"):
        p.add_argument("--"+key,default=DEFAULTS[key])
    p.add_argument("--state-mode",choices=["all","running"],default=DEFAULTS["state_mode"])
    p.add_argument("--power-csv",type=Path)
    for key in ("power_threshold","state_max_gap","min_running_fraction","min_interval","max_interval"):
        p.add_argument("--"+key.replace("_","-"),type=float,default=DEFAULTS[key])
    for key in ("min_samples","bins","chunk_size"):
        p.add_argument("--"+key.replace("_","-"),type=int,default=DEFAULTS[key])
    p.add_argument("--bin-scale",choices=["log","linear"],default=DEFAULTS["bin_scale"])
    p.add_argument("--bin-edges",help="递增正数分箱边界，单位秒，例如0.001,0.5,1.5,5,60,550,650,86400")
    p.add_argument("--output",type=Path,help="唯一输出SVG文件路径")
    p.add_argument("--overwrite",action="store_true",help="明确覆盖同名SVG")
    p.add_argument("--slices",type=int,default=DEFAULTS["slices"],help="绘制切片数；0表示所有窗口（默认），正数表示等间距取样")
    p.add_argument("--smooth-points",type=int,default=DEFAULTS["smooth_points"],help="每条曲线的平滑评估点数，默认2400")
    p.add_argument("--self-check",action="store_true")
    return p


def main():
    if hasattr(sys.stdout,"reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    args=parser().parse_args()
    if args.self_check:
        self_check(); return
    points,workbook=selected_points(args.unit)
    if args.list_sensors:
        for i,row in enumerate(points,1):
            print(f"{i:02d}  {row[0]}  {row[1]}")
        return
    parse_period(args.window);parse_period(args.step)
    if args.chunk_size<1 or args.min_samples<1 or not 2<=args.bins<=10000:
        raise ValueError("窗口、步长、分块和样本阈值须为正；bins须在2至10000。")
    if not 0<=args.min_running_fraction<=1 or not np.isfinite(args.state_max_gap) or args.state_max_gap<=0:
        raise ValueError("运行支持参数不合法。")
    if args.bin_edges:
        edges=np.array([float(x) for x in args.bin_edges.split(",")])
    else:
        if not 0<args.min_interval<args.max_interval or not np.isfinite(args.max_interval):
            raise ValueError("须满足0<min-interval<max-interval，且边界有限。")
        space=np.geomspace if args.bin_scale=="log" else np.linspace
        edges=space(args.min_interval,args.max_interval,args.bins+1)
    if len(edges)<3 or not np.isfinite(edges).all() or np.any(edges<=0) or np.any(np.diff(edges)<=0):
        raise ValueError("分箱边界须至少3个，严格递增且为有限正数。")
    source=(args.input or ROOT/"raw_data"/args.unit/"csv"/(args.sensor+".csv")).resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    sensor=source.stem
    name=args.name or next((str(r[1]) for r in points if str(r[0])==sensor),sensor)
    output=(args.output or Path(DEFAULTS["output_dir"])/(sensor+"_interval_probability.svg")).resolve()
    raw=(ROOT/"raw_data").resolve()
    if output.suffix.lower()!=".svg" or output==raw or raw in output.parents:
        raise ValueError("输出须为raw_data之外的.svg文件。")
    if output.exists() and not args.overwrite:
        raise FileExistsError("SVG已存在；换路径或使用--overwrite。")
    if not 0<=args.slices<=100:
        raise ValueError("slices须在0至100之间，0表示全部窗口。")
    if not 200<=args.smooth_points<=20000:
        raise ValueError("smooth-points须在200至20000之间。")
    first,last=source_extent(source)
    start=time_ns(args.start) if args.start else first
    end=time_ns(args.end) if args.end else last+1
    spans=None;power=None;power_rows=0
    before=[metadata(source)]
    if args.state_mode=="running":
        if args.input and not args.power_csv and source.parent!=(ROOT/"raw_data"/args.unit/"csv").resolve():
            raise ValueError("外部CSV运行筛选须明确指定--power-csv。")
        tag=next((str(r[0]) for r in points if len(r)>2 and str(r[2])=="active_power"),None)
        power=args.power_csv or (ROOT/"raw_data"/args.unit/"csv"/(tag+".csv") if tag else None)
        if power is None or not power.is_file():
            raise ValueError("运行筛选需要出力CSV；请指定--power-csv。")
        power=power.resolve();before.append(metadata(power))
        a,b,power_rows=power_spans(power,args.power_threshold,args.state_max_gap,args.chunk_size)
        if not len(a):
            raise ValueError("没有满足规则的运行支持区间。")
        spans=(a,b)
        if args.start is None:start=max(start,int(a[0]))
        if args.end is None:end=min(end,int(b[-1])+1)
    if end<=start:
        raise ValueError("end必须大于start。")
    starts,stops=make_windows(start,end,args.window,args.step,1_000_000//(len(edges)+3))
    n=len(starts)
    duration=(stops-starts)/SECOND
    coverage=duration if spans is None else support_before(stops,*spans)-support_before(starts,*spans)
    print(f"扫描：{name}；{n}个滑窗；唯一输出：{output.name}",flush=True)
    hist,audit=scan_intervals(source,start,end,None,None,edges,spans,args.chunk_size,(starts,stops))
    nb=len(edges)-1
    counts=hist[:,:nb+2].sum(axis=1)
    usable=(counts>=args.min_samples)&(coverage/duration>=args.min_running_fraction)
    all_prob=np.divide(hist[:,:nb+2],counts[:,None],out=np.full((n,nb+2),np.nan),where=counts[:,None]>0)
    np.testing.assert_allclose(all_prob[counts>0].sum(axis=1),1)
    probabilities=all_prob[:,1:-1].copy();probabilities[~usable]=np.nan
    after=[metadata(source)]+([metadata(power)] if power is not None else [])
    if before!=after:
        raise ValueError("扫描期间来源发生变化。")
    provenance=dict(
        sensor=sensor,name=name,timezone="Asia/Shanghai",
        analysed_start=str(pd.Timestamp(start)),analysed_end_exclusive=str(pd.Timestamp(end)),
        parameters={k:str(v) if isinstance(v,Path) else v for k,v in vars(args).items()},
        source_files=before,source_audit=audit,power_rows_scanned=power_rows,
        script_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        workbook=metadata(workbook) if workbook else None,
        environment=dict(numpy=np.__version__,pandas=pd.__version__,matplotlib=matplotlib.__version__,scipy=scipy.__version__),
        definition="P(interval in bin | interval starts in window)=bin_count/all eligible completed interval_count; probability mass, not density",
        window_rule="Interval left endpoint in window; right endpoint can follow window/analysis end. All windows retained; tail windows shortened.",
        quality_rule="Positive adjacent finite-valued records only. Duplicate timestamps add no zero interval; descending timestamps abort. Final unclosed interval censored.",
        state_rule="All periods by default; optional running mode requires endpoints in same continuously supported-running span.",
        bin_edges_seconds=edges.tolist(),
        normalization="All eligible intervals including underflow/overflow in denominator; last visible bin includes right boundary. No visible-range renormalization.",
        underflow_counts=hist[:,0].tolist(),overflow_counts=hist[:,nb+1].tolist(),
        excluded_state_crossing_counts=hist[:,nb+2].tolist(),right_censored_counts=hist[:,nb+3].tolist(),
        analysis_supported_seconds=coverage.tolist(),
        notes="No KDE, no per-second density, no theoretical fit; overlapping windows dependent; insufficient windows remain NaN."
    )
    output.parent.mkdir(parents=True,exist_ok=True)
    write_slices(output,starts,stops,edges,probabilities,hist[:,1:nb+1],counts,usable,name,args,provenance)
    print(f"完成：{output}\n全部窗口：{n}；可估计窗口：{int(usable.sum())}；仅生成这张SVG。")


if __name__=="__main__":
    main()
