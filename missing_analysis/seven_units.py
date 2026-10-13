"""Seven independent unit analyses, with only three SVG figures per K.

python missing_analysis/seven_units.py all --workers 3
python missing_analysis/seven_units.py report
"""
from __future__ import annotations
import argparse
from concurrent.futures import ProcessPoolExecutor, as_completed
import hashlib
import html
import importlib.metadata
import itertools
import json
import os
from pathlib import Path
import platform
import sys
import time
import zipfile

import numpy as np
import pandas as pd
from sklearn.metrics import adjusted_rand_score, silhouette_samples
from threadpoolctl import threadpool_limits

# Reuse the already validated numerical definitions; never call its old path-specific CLI.
from jqrd_gt1 import daily_modes as core

ROOT = Path(__file__).resolve().parents[1]
HERE = Path(__file__).resolve().parent
DEFAULT_OUT = HERE/'seven_units_v1'
UNITS = ['gard_gt1', 'gard_gt2', 'jqrd_gt1', 'jqrd_gt2', 'jxrd_gt1', 'jxrd_gt2', 'jxrd_gt3']
HISTORY = HERE/'jqrd_gt1'/'runs'/'jqrd_gt1_daily_v1'
PROTOCOL = dict(core.CONFIG, protocol='seven_units_daily_v1', unit='independent_per_unit',
                units=UNITS, retained_fuel_components=['CH4'],
                min_stability_checks=3, stable_agreement=.8,
                plots_per_k=['class_timeline.svg', 'representative_distributions.svg', 'transition_matrices.svg'],
                cross_unit_class_identity=False, shared_tag_handling='separate_unit_analyses_not_independent_evidence')


def cache_dir(out, unit):
    p = out/'_data'/unit
    p.mkdir(parents=True, exist_ok=True)
    return p


def save_npz(path, arrays):
    temp = path.with_suffix('.tmp')
    with temp.open('wb') as f:
        np.savez_compressed(f, **arrays)
    temp.replace(path)


def points(unit):
    from openpyxl import load_workbook
    files = list((ROOT/'raw_data'/unit).glob('*选取点表.xlsx'))
    assert len(files) == 1
    w = load_workbook(files[0], read_only=True, data_only=True)
    rows = [dict(sensor=r[0], name=r[1], semantic=r[2]) for r in
            w.worksheets[0].iter_rows(min_row=2, values_only=True) if r[0] and r[2] not in core.EXCLUDE]
    w.close()
    assert len(rows) == (30 if unit.startswith('gard') else 23)
    assert sum(r['semantic'] == 'CH4' for r in rows) == 1
    return rows, files[0]


def source_inventory(unit):
    rows, workbook = points(unit)
    sources = []
    for row in rows:
        path = ROOT/'raw_data'/unit/'csv'/f'{row["sensor"]}.csv'
        first, last = core.source_extent(path)
        sources.append(dict(**row, **core.metadata(path), first_ns=first, last_ns=last))
    return dict(workbook=core.metadata(workbook), workbook_sha256=core.digest(workbook), sources=sources)


def check_sources(inventory):
    for s in inventory['sources']:
        assert core.metadata(Path(s['path'])) == {k:s[k] for k in ['path','bytes','mtime_ns']}, s['path']
    w = inventory['workbook']
    assert core.metadata(Path(w['path'])) == w
    assert core.digest(Path(w['path'])) == inventory['workbook_sha256']


def history_inventory_matches(inventory):
    path=HISTORY/'sources.json'
    return path.exists() and core.read_json(path)==inventory


def history_models_ready(cache):
    audit=cache/'audit.json'
    if not audit.exists() or core.read_json(audit).get('reused_from')!=str(HISTORY):
        return False
    required=['baseline.npz','daily_labels.csv','initializations.json','class_stability.csv',
              'block_stability.csv','initialization_ari.csv','evaluation.json']
    required += [f'perturbation_{rep:02d}{suffix}' for rep in range(20)
                 for suffix in ['.json','_checks.npz']]
    return all((HISTORY/f'k{k}'/name).exists() for k in PROTOCOL['k_values'] for name in required)


def inventory_run(out):
    out.mkdir(parents=True, exist_ok=True)
    path = out/'_data'/'manifest.json'
    inventory = {u:source_inventory(u) for u in UNITS}
    if path.exists():
        old = core.read_json(path)
        assert old['protocol'] == PROTOCOL, 'Protocol changed: use a new output directory'
        assert old['inventory'] == inventory, 'Source changed: use a new output directory'
        return old
    tag_units = {}
    for u, item in inventory.items():
        for s in item['sources']:
            tag_units.setdefault(s['sensor'], []).append(u)
    manifest = dict(protocol=PROTOCOL, inventory=inventory,
        possible_shared_tags={tag:units for tag,units in tag_units.items() if len(units)>1},
        created_utc=pd.Timestamp.now(tz='UTC').isoformat(), python=sys.version, platform=platform.platform(),
        dependencies={p:importlib.metadata.version(p) for p in
                      ['numpy','pandas','scipy','scikit-learn','kmedoids','matplotlib','openpyxl','threadpoolctl']})
    assert manifest['dependencies']['kmedoids'] == '0.5.5'
    core.save_json(path, manifest)
    return manifest


def make_daily_frame(s, order, origin, n_days, z):
    d = np.arange(n_days)
    first, last = (s['first_ns']-origin)//core.DAY, (s['last_ns']-origin)//core.DAY
    covered = (d >= first) & (d <= last)
    boundary = (d == first) | (d == last)
    total = z['counts'].sum(axis=1)
    status = np.full(n_days, 'eligible', dtype=object)
    status[total < 50] = 'insufficient_intervals'
    status[z['observations'] == 0] = 'no_valid_observations'
    status[~covered] = 'outside_coverage'
    status[boundary & (total >= 50)] = 'boundary_assignable'
    f = pd.DataFrame(dict(sensor=s['sensor'], name=s['name'], semantic=s['semantic'], sensor_order=order,
        date=pd.date_range(pd.Timestamp(origin),periods=n_days,freq='D').strftime('%Y-%m-%d'),
        day_index=d, status=status, boundary=boundary, valid_observations=z['observations'],
        completed_intervals=total, longest_interval_seconds=z['longest'],
        long_gap_overlap_seconds=z['long_gap_seconds'], unfinished_intervals=z['unfinished'],
        assignable=covered & (total >= 50), train_eligible=covered & ~boundary & (total >= 50)))
    for name, pos in [('underflow_probability',0),('overflow_probability',51)]:
        f[name] = np.divide(z['counts'][:,pos],total,out=np.full(n_days,np.nan),where=total>0)
    return f


def prepare_unit(out, unit, inventory):
    c = cache_dir(out, unit)
    check_sources(inventory)
    if (c/'audit.json').exists() and (c/'daily.npz').exists() and (c/'daily.csv').exists():
        print(f'{unit}: prepared cache verified', flush=True)
        return
    # Reuse JQRD GT1 only when the complete source inventory is identical.
    if unit == 'jqrd_gt1' and history_inventory_matches(inventory):
        f = pd.read_csv(HISTORY/'daily_metadata.csv')
        z = dict(np.load(HISTORY/'daily_profiles.npz'))
        core.csv(c/'daily.csv', f)
        save_npz(c/'daily.npz',z)
        old_audit = core.read_json(HISTORY/'source_audit.json')
        core.save_json(c/'audit.json',dict(sources=old_audit,reused_from=str(HISTORY),
                       rows=sum(a['rows'] for a in old_audit),positive_intervals=int(z['counts'].sum()),
                       source_metadata_unchanged=True))
        print(f'{unit}: reused verified historical daily profiles',flush=True)
        return
    sources = inventory['sources']
    origin = min(s['first_ns'] for s in sources)//core.DAY*core.DAY
    last = max(s['last_ns'] for s in sources)//core.DAY*core.DAY
    n_days = int((last-origin)//core.DAY+1)
    partial = c/'scan_checkpoint.npz'
    saved = dict(np.load(partial)) if partial.exists() else {}
    frames, hist, audits = [], [], []
    for i,s in enumerate(sources):
        prefix = f's{i}_'
        if prefix+'audit' in saved:
            audit = json.loads(str(saved[prefix+'audit']))
            z = {key:saved[prefix+key] for key in ['counts','observations','longest','long_gap_seconds','unfinished']}
        else:
            print(f'{unit}: scan {i+1}/{len(sources)} {s["sensor"]} ({s["bytes"]/2**30:.2f} GiB)',flush=True)
            z,audit = core.scan_source(Path(s['path']),origin,n_days,PROTOCOL['chunk_size'])
            assert core.metadata(Path(s['path'])) == {k:s[k] for k in ['path','bytes','mtime_ns']}
            for key,value in z.items():
                saved[prefix+key] = value
            saved[prefix+'audit'] = np.array(json.dumps(core.clean(audit)))
            save_npz(partial,saved)
        assert audit['first_ns']==s['first_ns'] and audit['last_ns']==s['last_ns']
        assert z['counts'].sum()==audit['positive_intervals']
        frames.append(make_daily_frame(s,i,origin,n_days,z))
        hist.append(z['counts'])
        audits.append(dict(sensor=s['sensor'],**audit))
    f = pd.concat(frames,ignore_index=True)
    f.insert(0,'sample_id',np.arange(len(f)))
    counts=np.vstack(hist)
    totals=counts.sum(axis=1)
    p=np.divide(counts,totals[:,None],out=np.full(counts.shape,np.nan),where=totals[:,None]>0)
    np.testing.assert_allclose(p[totals>0].sum(axis=1),1)
    assert int(counts.sum())==sum(a['positive_intervals'] for a in audits)
    core.csv(c/'daily.csv',f)
    save_npz(c/'daily.npz',dict(counts=counts,probabilities=p,edges=core.EDGES))
    check_sources(inventory)
    core.save_json(c/'audit.json',dict(sources=audits,rows=sum(a['rows'] for a in audits),
                   positive_intervals=int(counts.sum()),source_metadata_unchanged=True))
    # Delete only our redundant, specifically named checkpoint after validated finalization.
    partial.unlink(missing_ok=True)
    print(f'{unit}: prepared {len(f)} sensor-days, {int(f.train_eligible.sum())} learning samples',flush=True)


def matrices(labels, stable, f, k):
    all_counts=np.zeros((k,k),dtype=np.int64)
    stable_counts=all_counts.copy()
    for _,group in f.groupby('sensor',sort=False):
        ids=group.index.to_numpy()
        all_counts+=core.transitions(labels[ids],k)
        stable_counts+=core.transitions(np.where(stable[ids],labels[ids],-1),k)
    return all_counts,stable_counts


def class_stats(f,p,train,medoids,labels,silhouette,class_j):
    rows=[]
    bands=[]
    for c,m in enumerate(medoids):
        ids=train[labels[train]==c]
        bands.append(np.quantile(p[ids],[.1,.9],axis=0))
        row=f.iloc[m]
        rows.append(dict(cluster=c+1,learning_days=len(ids),assigned_days=int((labels==c).sum()),
            sensor_count=f.iloc[ids].sensor.nunique(),distinct_dates=f.iloc[ids].date.nunique(),
            medoid_sample_id=int(m),sensor=row.sensor,name=row['name'],date=row.date,
            median_bin=int(np.searchsorted(np.cumsum(p[m]),.5)),
            silhouette_mean=float(silhouette[labels[train]==c].mean()),
            jaccard_median=core.quantile(class_j[:,c],.5),jaccard_p10=core.quantile(class_j[:,c],.1)))
    return rows,np.asarray(bands)


def import_history(unit,f,p):
    assert unit=='jqrd_gt1'
    models={}
    summaries={}
    train=np.flatnonzero(f.train_eligible)
    assign=np.flatnonzero(f.assignable)
    models.update(train_ids=train,assign_ids=assign)
    for k in PROTOCOL['k_values']:
        root=HISTORY/f'k{k}'
        base=dict(np.load(root/'baseline.npz'))
        labels=base['labels']
        table=pd.read_csv(root/'daily_labels.csv')
        checks=table.check_count.to_numpy()
        fraction=table.agreement_fraction.to_numpy()
        stable=table.stability.eq('stable').to_numpy()
        codes=np.full(len(f),-1,dtype=np.int8)
        codes[assign]=0 # 0 insufficient evidence; 1 unstable; 2 stable
        codes[table.stability.eq('unstable')]=1
        codes[stable]=2
        init=core.read_json(root/'initializations.json')
        class_table=pd.read_csv(root/'class_stability.csv')
        class_j=class_table.jaccard.to_numpy().reshape(20,k)
        stats,bands=class_stats(f,p,train,base['medoid_ids'],labels,base['silhouette'],class_j)
        all_t,stable_t=matrices(labels,stable,f,k)
        data=dict(labels=labels,medoids=base['medoid_ids'],silhouette=base['silhouette'],
            check_count=checks,agreement=fraction,stability=codes,class_jaccard=class_j,
            block_ari=pd.read_csv(root/'block_stability.csv').ari.to_numpy(),
            init_ari=pd.read_csv(root/'initialization_ari.csv').ari.to_numpy(),
            init_labels=np.array([x['labels'] for x in init]),init_medoids=np.array([x['medoid_sample_ids'] for x in init]),
            init_losses=np.array([x['loss'] for x in init]),init_seeds=np.array([x['seed'] for x in init]),
            all_transitions=all_t,stable_transitions=stable_t,bands=bands)
        omitted=np.full((20,len(f)),-1,dtype=np.int16)
        seeds=[];losses=[];medoid_runs=[];mapping=[]
        missing=[]
        for rep in range(20):
            z=np.load(root/f'perturbation_{rep:02d}_checks.npz')
            omitted[rep,z['sample_ids']]=z['aligned']
            obj=core.read_json(root/f'perturbation_{rep:02d}.json')
            seeds.append([r['seed'] for r in obj['initializations']])
            losses.append([r['loss'] for r in obj['initializations']])
            medoid_runs.append([r['medoid_sample_ids'] for r in obj['initializations']])
            mapping.append(obj['mapping'])
            q=class_table[class_table.repetition==rep]
            missing.append(q[['baseline_absent','resampled_absent','train_baseline_absent']].to_numpy())
        data.update(omitted_labels=omitted,perturbation_seeds=np.array(seeds),perturbation_losses=np.array(losses),
                    perturbation_medoids=np.array(medoid_runs),perturbation_mapping=np.array(mapping),
                    missing_classes=np.array(missing))
        models.update({f'k{k}_{key}':value for key,value in data.items()})
        summary=core.read_json(root/'evaluation.json')
        summary.update(classes=stats,reused_historical_model=True)
        summaries[str(k)]=summary
    return models,summaries


def evaluate_unit(out,unit):
    threadpool_limits(limits=1)
    c=cache_dir(out,unit)
    manifest=core.read_json(out/'_data'/'manifest.json')
    check_sources(manifest['inventory'][unit])
    f=pd.read_csv(c/'daily.csv')
    p=np.load(c/'daily.npz')['probabilities']
    model_path=c/'models.npz'
    summary_path=c/'analysis.json'
    if unit=='jqrd_gt1' and not summary_path.exists() and history_models_ready(c):
        models,summaries=import_history(unit,f,p)
        save_npz(model_path,models)
        core.save_json(summary_path,dict(k_results=summaries))
        print(f'{unit}: reused all 350 historical fits and date-block checks',flush=True)
    models=dict(np.load(model_path)) if model_path.exists() else {}
    summary=core.read_json(summary_path) if summary_path.exists() else dict(k_results={})
    train=np.flatnonzero(f.train_eligible)
    assign=np.flatnonzero(f.assignable)
    n_days=int(f.day_index.max()+1)
    models.update(train_ids=train,assign_ids=assign)
    models['retained_date_masks']=np.array([core.date_split(n_days,r)[0] for r in range(20)])
    need_fit=any(str(k) not in summary['k_results'] for k in PROTOCOL['k_values'])
    # Temporary matrix is removed at context exit; profiles and all models stay cached.
    import tempfile
    with tempfile.TemporaryDirectory(prefix=f'{unit}_distance_',dir=c) as td:
        distance=None
        if need_fit:
            path=Path(td)/'distance.npy'
            distance=np.lib.format.open_memmap(path,mode='w+',dtype='float64',shape=(len(train),len(train)))
            for start in range(0,len(train),256):
                distance[start:start+256]=core.hellinger(p[train[start:start+256]],p[train])
            np.fill_diagonal(distance,0)
            distance.flush()
            for start in range(0,len(train),512):
                np.testing.assert_allclose(distance[start:start+512],distance[:,start:start+512].T,atol=1e-14)
        for k in PROTOCOL['k_values']:
            if str(k) in summary['k_results']:
                continue
            print(f'{unit} K={k}: 10 starts, then 20 date-block repetitions',flush=True)
            starts=[core.fit_once(distance,k,core.seeded(10,k,i)) for i in range(10)]
            best=min(range(10),key=lambda i:starts[i]['loss'])
            medoids=core.order_medoids(train[starts[best]['medoids']],p,f)
            da=core.hellinger(p[assign],p[medoids])
            labels=np.full(len(f),-1,dtype=np.int16)
            labels[assign]=np.argmin(da,axis=1)
            silhouette=silhouette_samples(distance,labels[train],metric='precomputed')
            ari_init=np.array([adjusted_rand_score(a['labels'],b['labels']) for a,b in itertools.combinations(starts,2)])
            checks=np.zeros(len(f),dtype=int); agrees=checks.copy()
            omitted=np.full((20,len(f)),-1,dtype=np.int16)
            block_ari=[];class_j=[];seeds=[];losses=[];medoid_runs=[];mappings=[];missing=[]
            for rep in range(20):
                keep=models['retained_date_masks'][rep]
                retained=np.flatnonzero(keep[f.day_index.to_numpy()[train]])
                test_ids=assign[~keep[f.day_index.to_numpy()[assign]]]
                assert not set(f.iloc[train[retained]].day_index)&set(f.iloc[test_ids].day_index)
                sub=np.ascontiguousarray(distance[np.ix_(retained,retained)])
                fits=[core.fit_once(sub,k,core.seeded(30,k,rep,i)) for i in range(3)]
                fit=min(fits,key=lambda x:x['loss'])
                new_medoids=train[retained[fit['medoids']]]
                mapping,overlap=core.match_labels(labels[train[retained]],fit['labels'],k)
                predicted=np.argmin(core.hellinger(p[test_ids],p[new_medoids]),axis=1)
                aligned=mapping[predicted]
                omitted[rep,test_ids]=aligned
                checks[test_ids]+=1;agrees[test_ids]+=aligned==labels[test_ids]
                block_ari.append(adjusted_rand_score(labels[test_ids],predicted))
                scores=[];missing_rep=[]
                for j in range(k):
                    a=labels[test_ids]==j;b=aligned==j
                    union=(a|b).sum()
                    scores.append((a&b).sum()/union if union else np.nan)
                    missing_rep.append([not a.any(),not b.any(),not (labels[train[retained]]==j).any()])
                class_j.append(scores);missing.append(missing_rep)
                seeds.append([x['seed'] for x in fits]);losses.append([x['loss'] for x in fits])
                medoid_runs.append([train[retained[x['medoids']]] for x in fits]);mappings.append(mapping)
                if rep%5==4:
                    print(f'{unit} K={k}: date blocks {rep+1}/20',flush=True)
                del sub
            agreement=np.divide(agrees,checks,out=np.full(len(f),np.nan),where=checks>0)
            stability=np.full(len(f),-1,dtype=np.int8)
            stability[assign]=0
            stability[(checks>=3)&(agreement<.8)]=1
            stability[(checks>=3)&(agreement>=.8)]=2
            class_j=np.array(class_j)
            stats,bands=class_stats(f,p,train,medoids,labels,silhouette,class_j)
            all_t,stable_t=matrices(labels,stability==2,f,k)
            result=dict(k=k,best_initialization=best,silhouette_mean=float(silhouette.mean()),
                mean_hellinger=float(distance[np.arange(len(train)),np.searchsorted(train,medoids)[labels[train]]].mean()),
                init_ari_median=core.quantile(ari_init,.5),init_ari_p10=core.quantile(ari_init,.1),
                block_ari_median=core.quantile(block_ari,.5),block_ari_p10=core.quantile(block_ari,.1),
                min_class_jaccard_median=float(np.min([s['jaccard_median'] for s in stats])),
                unstable_days=int((stability==1).sum()),insufficient_evidence_days=int((stability==0).sum()))
            t=PROTOCOL['thresholds']
            gates=dict(init_ari=result['init_ari_median']>=t['init_ari'],
                 block_ari=result['block_ari_median']>=t['block_ari'],
                 block_ari_p10=result['block_ari_p10']>=t['block_ari_p10'],
                 class_jaccard=result['min_class_jaccard_median']>=t['class_jaccard'],
                 silhouette=result['silhouette_mean']>0)
            result.update(passes=all(gates.values()),gates=gates,classes=stats)
            data=dict(labels=labels,medoids=medoids,silhouette=silhouette,check_count=checks,agreement=agreement,
                stability=stability,class_jaccard=class_j,block_ari=np.array(block_ari),init_ari=ari_init,
                init_labels=np.array([s['labels'] for s in starts]),init_medoids=np.array([train[s['medoids']] for s in starts]),
                init_losses=np.array([s['loss'] for s in starts]),init_seeds=np.array([s['seed'] for s in starts]),
                all_transitions=all_t,stable_transitions=stable_t,bands=bands,omitted_labels=omitted,
                perturbation_seeds=np.array(seeds),perturbation_losses=np.array(losses),
                perturbation_medoids=np.array(medoid_runs),perturbation_mapping=np.array(mappings),
                missing_classes=np.array(missing))
            # Exact replay of the selected initialization on the actual unit data.
            replay=core.fit_once(distance,k,starts[best]['seed'])
            np.testing.assert_array_equal(replay['medoids'],starts[best]['medoids'])
            np.testing.assert_array_equal(replay['labels'],starts[best]['labels'])
            result['seed_replay_passed']=True
            models.update({f'k{k}_{key}':value for key,value in data.items()})
            summary['k_results'][str(k)]=result
            save_npz(model_path,models)
            core.save_json(summary_path,summary)
            print(f'{unit} K={k}: silhouette={result["silhouette_mean"]:.4f}, block ARI={result["block_ari_median"]:.4f}, passes={result["passes"]}',flush=True)
        if distance is not None:
            distance.flush()
            distance._mmap.close()
            del distance
    eligible=[s for s in summary['k_results'].values() if s['passes']]
    recommended=None
    if eligible:
        highest=max(s['silhouette_mean'] for s in eligible)
        recommended=min(s['k'] for s in eligible if s['silhouette_mean']>=highest-.02)
    summary.update(recommended_k=recommended,confirmed_k=5 if unit=='jqrd_gt1' else None)
    save_npz(model_path,models)
    core.save_json(summary_path,summary)
    check_sources(manifest['inventory'][unit])


def temporal_summary(f,models,k):
    labels=models[f'k{k}_labels']
    sensor_rows=[];month_rows=[];runs=[];nulls=[]
    for si,(sensor,g) in enumerate(f.groupby('sensor',sort=False)):
        ids=g.index.to_numpy()
        lab=labels[ids]
        observed=core.temporal_metrics(lab,k)
        rng=np.random.default_rng(core.seeded(40,k,si))
        valid=np.flatnonzero(lab>=0)
        shuffled=lab.copy()
        null=np.full((500,3),np.nan)
        for rep in range(500):
            shuffled[valid]=rng.permutation(lab[valid])
            null[rep]=core.temporal_metrics(shuffled,k)
        nulls.append(null)
        pair=core.transitions(lab,k)
        sensor_rows.append(dict(sensor=sensor,name=g.iloc[0]['name'],classified_days=len(valid),
            proportions=[float((lab==j).sum()/len(valid)) if len(valid) else None for j in range(k)],
            adjacent_pairs=int(pair.sum()),switches=int(pair.sum()-np.trace(pair)),
            self_transition=observed[0],mean_visible_run=observed[1],max_visible_run=observed[2],
            null_median=[core.quantile(null[:,j],.5) for j in range(3)],
            null_p975=[core.quantile(null[:,j],.975) for j in range(3)],
            p_greater=[(1+int((null[:,j]>=observed[j]).sum()))/501 if np.isfinite(observed[j]) else None for j in range(3)]))
        for row in core.visible_runs(lab,g.boundary.to_numpy()):
            runs.append(dict(sensor_order=si,**row))
        months=g.date.str[:7].to_numpy()
        for month in np.unique(months):
            sel=months==month
            valid_pairs=(lab[:-1]>=0)&(lab[1:]>=0)&sel[1:]
            month_rows.append(dict(sensor_order=si,month=month,classified_days=int((lab[sel]>=0).sum()),
                class_counts=[int((lab[sel]==j).sum()) for j in range(k)],
                adjacent_pairs=int(valid_pairs.sum()),switches=int(((lab[1:]!=lab[:-1])&valid_pairs).sum())))
    return dict(sensors=sensor_rows,months=month_rows,visible_runs=runs),np.array(nulls)


def plot_three(out,unit,f,p,models,k,stats):
    from matplotlib.colors import ListedColormap, BoundaryNorm
    from matplotlib.patches import Patch
    plt=core.plot_setup()
    folder=out/unit/f'k{k}'
    folder.mkdir(parents=True,exist_ok=True)
    labels=models[f'k{k}_labels']
    medoids=models[f'k{k}_medoids']
    bands=models[f'k{k}_bands']
    n_sensors=f.sensor.nunique()
    n_days=int(f.day_index.max()+1)
    colors=list(plt.get_cmap('tab10').colors[:k])
    data=np.where(f.status.eq('outside_coverage'),-3,np.where(f.status.eq('no_valid_observations'),-2,-1))
    data[labels>=0]=labels[labels>=0]
    fig,ax=plt.subplots(figsize=(17,max(9,n_sensors*.38)),layout='constrained')
    palette=['#f4f4f4','#9e9e9e','#d4d4d4']+colors
    ax.imshow(data.reshape(n_sensors,n_days),aspect='auto',interpolation='nearest',
        cmap=ListedColormap(palette),norm=BoundaryNorm(np.arange(-3.5,k+.5),len(palette)))
    names=f.drop_duplicates('sensor')['name'].tolist()
    ax.set_yticks(np.arange(n_sensors),names)
    dates=f[f.sensor_order==0].date.to_numpy()
    ticks=np.r_[0,[i for i in range(1,n_days) if dates[i][8:]=='01']]
    if len(ticks)>18:
        ticks=ticks[::2]
    ax.set_xticks(ticks,[dates[i][:7] for i in ticks],rotation=40,ha='right')
    ax.set_xlabel('上海时间 / 日期')
    ax.set_title(f'{unit.upper().replace("_", " ")} 每日记录间隔模式 | K={k}')
    legends=['覆盖外','无有效记录','间隔不足']+[f'C{j+1}' for j in range(k)]
    ax.legend(handles=[Patch(color=c,label=l) for c,l in zip(palette,legends)],
              ncol=k+3,loc='upper center',bbox_to_anchor=(.5,1.10),fontsize=9)
    core.save_svg(fig,folder/'class_timeline.svg')
    fig,axes=plt.subplots(k,1,figsize=(12,2.65*k),sharex=True,layout='constrained')
    for j,(ax,m) in enumerate(zip(axes,medoids)):
        row=f.iloc[m]
        ax.fill_between(np.arange(52),bands[j,0],bands[j,1],step='mid',color=colors[j],alpha=.2,label='类内10%—90%')
        ax.plot(p[m],color=colors[j],drawstyle='steps-mid',lw=1.6,label='真实代表日')
        ax.set_title(f'{unit.upper()} C{j+1} | {row.sensor} | {row["name"]} | {row.date}')
        core.distribution_axis(ax)
        ax.legend(loc='upper right')
    core.save_svg(fig,folder/'representative_distributions.svg')
    fig,axes=plt.subplots(1,2,figsize=(12,5.3),layout='constrained')
    for ax,scope,title in zip(axes,['all','stable'],['全部已分类日期','仅稳定日期（其余断开）']):
        count=models[f'k{k}_{scope}_transitions']
        den=count.sum(axis=1)
        prob=np.divide(count,den[:,None],out=np.full((k,k),np.nan),where=den[:,None]>0)
        im=ax.imshow(prob,vmin=0,vmax=1,cmap='Blues')
        for a in range(k):
            for b in range(k):
                ax.text(b,a,f'{prob[a,b]:.2f}' if np.isfinite(prob[a,b]) else '—',ha='center',va='center',
                        color='white' if prob[a,b]>.5 else 'black',fontsize=9)
        ax.set_xticks(range(k),[f'C{j+1}' for j in range(k)])
        ax.set_yticks(range(k),[f'C{j+1}' for j in range(k)])
        ax.set_xlabel('次日类别');ax.set_ylabel('当日类别')
        ax.set_title(f'{title} | 相邻日对 {count.sum():,}')
        fig.colorbar(im,ax=ax,shrink=.8)
    fig.suptitle(f'{unit.upper().replace("_", " ")} | K={k}')
    core.save_svg(fig,folder/'transition_matrices.svg')


def report_unit(out,unit):
    threadpool_limits(limits=1)
    c=cache_dir(out,unit)
    f=pd.read_csv(c/'daily.csv')
    p=np.load(c/'daily.npz')['probabilities']
    models=dict(np.load(c/'models.npz'))
    summary=core.read_json(c/'analysis.json')
    for k in PROTOCOL['k_values']:
        s=summary['k_results'][str(k)]
        if 'temporal' not in s:
            s['temporal'],null=temporal_summary(f,models,k)
            models[f'k{k}_permutation_values']=null
        plot_three(out,unit,f,p,models,k,s)
        print(f'{unit} K={k}: 3 SVGs complete',flush=True)
    for k in PROTOCOL['k_values']:
        lab=models[f'k{k}_labels']
        f[f'k{k}_class']=np.where(lab>=0,lab+1,np.nan)
        f[f'k{k}_stability']=models[f'k{k}_stability']
        f[f'k{k}_check_count']=models[f'k{k}_check_count']
        f[f'k{k}_agreement']=models[f'k{k}_agreement']
    core.csv(c/'daily.csv',f)
    save_npz(c/'models.npz',models)
    summary['status_counts']={str(key):int(value) for key,value in f.status.value_counts().items()}
    summary['start']=str(f.date.min());summary['end']=str(f.date.max())
    summary['sensors']=int(f.sensor.nunique())
    summary['learning_days']=int(f.train_eligible.sum());summary['assigned_days']=int(f.assignable.sum())
    summary['report_complete']=True
    core.save_json(c/'analysis.json',summary)


def summarize(out):
    rows=[];sections=[];text=[]
    manifest=core.read_json(out/'_data'/'manifest.json')
    total_rows=0
    for unit in UNITS:
        c=cache_dir(out,unit)
        if not (c/'analysis.json').exists():
            continue
        summary=core.read_json(c/'analysis.json')
        if not summary.get('report_complete'):
            continue
        audit=core.read_json(c/'audit.json')
        total_rows+=audit['rows']
        recommended=summary['recommended_k']
        title=f'{unit.upper()} — {summary["sensors"]}测点；规则推荐K={recommended if recommended else "无合格方案"}'
        if summary['confirmed_k']:
            title+=f'；人工确认K={summary["confirmed_k"]}'
        body=[f'<section><h2>{html.escape(title)}</h2>',
              f'<p>{summary["start"]} 至 {summary["end"]}；学习日分布 {summary["learning_days"]:,}；可分类日 {summary["assigned_days"]:,}。</p>',
              '<table><tr><th>K</th><th>轮廓系数</th><th>日期扰动ARI<br>中位数 / 10%分位</th><th>全部门槛</th><th>查看三张图</th></tr>']
        text.append(title)
        text.append(f'  完整来源行数 {audit["rows"]:,}；日期 {summary["start"]} 至 {summary["end"]}；学习日 {summary["learning_days"]}。')
        text.append('  日期状态：'+json.dumps(summary['status_counts'],ensure_ascii=False))
        for k in PROTOCOL['k_values']:
            s=summary['k_results'][str(k)]
            row=dict(unit=unit,k=k,**{key:s[key] for key in ['silhouette_mean','mean_hellinger','init_ari_median',
                'init_ari_p10','block_ari_median','block_ari_p10','min_class_jaccard_median','passes','unstable_days',
                'insufficient_evidence_days']})
            row['recommended_k']=recommended;row['confirmed_k']=summary['confirmed_k']
            rows.append(row)
            links=' · '.join(f'<a target="_blank" href="{unit}/k{k}/{file}">{name}</a>' for file,name in
                [('class_timeline.svg','时间条带'),('representative_distributions.svg','代表分布'),('transition_matrices.svg','转移矩阵')])
            gate_names=dict(init_ari='初始化稳定性',block_ari='日期扰动中位数',
                            block_ari_p10='日期扰动10%分位',class_jaccard='逐类一致性',silhouette='轮廓系数')
            failed='、'.join(gate_names[key] for key,passed in s['gates'].items() if not passed)
            verdict='通过' if s['passes'] else '未通过：'+failed
            body.append(f'<tr><td>{k}</td><td>{s["silhouette_mean"]:.4f}</td><td>{s["block_ari_median"]:.4f} / {s["block_ari_p10"]:.4f}</td>'
                        f'<td>{verdict}</td><td>{links}</td></tr>')
            text.append(f'  K={k}: silhouette={s["silhouette_mean"]:.4f}, date ARI={s["block_ari_median"]:.4f}, '
                        f'p10={s["block_ari_p10"]:.4f}, all_gates={s["passes"]}。')
            if failed:
                text.append('    未通过：'+failed)
        five=summary['k_results']['5']
        models=np.load(c/'models.npz')
        count=models['k5_all_transitions']
        self_ratio=np.trace(count)/count.sum() if count.sum() else np.nan
        five_null=five['temporal']['sensors']
        persistent=sum(s['self_transition'] is not None and s['null_p975'][0] is not None and s['self_transition']>s['null_p975'][0] for s in five_null)
        description=f'K=5自转移比例 {self_ratio:.2%}；{persistent}/{summary["sensors"]}测点高于各自打乱对照97.5%分位数。'
        text.append('  '+description)
        body.append('</table><p>'+description+'</p></section>')
        sections.append('\n'.join(body))
    core.csv(out/'类别数比较.csv',rows)
    intro=['七台燃机每日记录间隔分析',
      '每台机组分别聚类；同一机组内所有选取测点共享分类。不同机组/不同K的C1等编号和颜色不表示同一分布。',
      '蓝色表示本模型排序后的C1，不能跨机组自动认定为完整高频真值或作为掩码保护区。',
      '燃料组分仅CH4，密度及高/低热值保留；GARD各30测点，其余各23测点。',
      '自然日、全时段、50个0.01—1000秒对数箱加上下溢；每日至少50间隔，首末边界日不参与学习。',
      'K=2—6，10初始化，20轮共同7日块80%子抽样，每轮3初始化；完整计算500次逐测点标签打乱。',
      '缺失原因及原始采集周期未识别，因此报告记录间隔与无记录日，不把所有秒级空位称为故障缺失。',
      '转移仅连接同测点相邻自然日；稳定日期图排除不稳定和证据不足日。连续可见段保留两端删失。',
      '所有数值结果集中在_data，公开查看入口为查看结果.html；图均为Matplotlib SVG，HTML仅提供本地链接。',
      'GARD/JQRD/JXRD存在同厂同编码的候选共享测点；逐机组结果不能作为独立样本重复累计证据。',
      f'完整扫描或验证缓存覆盖的文件行数合计 {total_rows:,}（包含不同目录中的潜在共享记录，不是独立观测总数）。',
      'JQRD GT1经完整来源清单一致性验证后复用历史缓存与全部候选模型；其他六台完整扫描。',
      '来源校验采用大小和纳秒修改时间，选取表另用SHA256；未对所有CSV另做完整内容哈希。',
      '稳定性和打乱是内部探索性证据；未进行独立外部验证，逐测点p值未经多重比较校正。','']
    (out/'分析说明.txt').write_text('\n'.join(intro+text)+'\n',encoding='utf-8')
    page='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>七台燃机记录间隔分析</title>
<style>body{font:16px/1.7 "Microsoft YaHei",sans-serif;max-width:1180px;margin:40px auto;padding:0 24px;background:#f6f8fa;color:#152536}h1{font-size:29px}h2{font-size:21px}section{background:white;padding:18px 24px;margin:24px 0;border-radius:10px}table{border-collapse:collapse;width:100%}th,td{padding:9px 12px;border-bottom:1px solid #ddd;text-align:left}a{color:#126bb5}p{color:#475569}</style>
<h1>七台燃机每日记录间隔分析</h1><p>每行提供三张SVG。按机组分别聚类，跨机组的同色类别不能直接等同；蓝色也不自动代表高质量真值。图中参数顺序沿用各自选取点表。</p>
<p><a href="类别数比较.csv">下载类别数比较</a> · <a href="分析说明.txt">查看分析说明</a></p>'''
    page+='\n'.join(sections)+'</html>'
    (out/'查看结果.html').write_text(page,encoding='utf-8')


def verify_outputs(out):
    manifest=core.read_json(out/'_data'/'manifest.json')
    result={}
    for unit in UNITS:
        check_sources(manifest['inventory'][unit])
        c=cache_dir(out,unit)
        f=pd.read_csv(c/'daily.csv')
        z=np.load(c/'daily.npz')
        models=np.load(c/'models.npz')
        audit=core.read_json(c/'audit.json')
        np.testing.assert_array_equal(z['counts'].sum(axis=1),f.completed_intervals)
        assert int(z['counts'].sum())==audit['positive_intervals']
        np.testing.assert_allclose(z['probabilities'][f.assignable].sum(axis=1),1)
        total=f.valid_observations.sum()
        duplicate=sum(s['duplicate_valid_timestamps'] for s in audit['sources'])
        n_valid_sources=sum(s['unclosed_last_valid_ns'] is not None for s in audit['sources'])
        assert total-duplicate-n_valid_sources==audit['positive_intervals']
        for k in PROTOCOL['k_values']:
            assert len(list((out/unit/f'k{k}').glob('*.svg')))==3
            m=models[f'k{k}_medoids'];lab=models[f'k{k}_labels']
            assert set(m).issubset(set(models['train_ids']))
            assert np.all(lab[~f.assignable]==-1)
            np.testing.assert_array_equal(np.argmin(core.hellinger(z['probabilities'][f.assignable],z['probabilities'][m]),axis=1),lab[f.assignable])
            held=models[f'k{k}_omitted_labels']
            np.testing.assert_array_equal((held>=0).sum(axis=0),models[f'k{k}_check_count'])
            train=models['train_ids'];assign=models['assign_ids']
            for rep in range(20):
                keep=models['retained_date_masks'][rep]
                checked=np.flatnonzero(held[rep]>=0)
                np.testing.assert_array_equal(checked,assign[~keep[f.iloc[assign].day_index]])
                ids=models[f'k{k}_perturbation_medoids'][rep]
                assert np.all(keep[f.iloc[ids.ravel()].day_index])
                # Reconstruct ALL omitted predictions from saved medoids and retained-only mappings.
                best=int(np.argmin(models[f'k{k}_perturbation_losses'][rep]))
                medoids=ids[best]
                retained=train[keep[f.iloc[train].day_index]]
                retained_labels=np.argmin(core.hellinger(z['probabilities'][retained],z['probabilities'][medoids]),axis=1)
                mapping,_=core.match_labels(lab[retained],retained_labels,k)
                np.testing.assert_array_equal(mapping,models[f'k{k}_perturbation_mapping'][rep])
                predicted=mapping[np.argmin(core.hellinger(z['probabilities'][checked],z['probabilities'][medoids]),axis=1)]
                np.testing.assert_array_equal(predicted,held[rep,checked])
            a,b=matrices(lab,models[f'k{k}_stability']==2,f,k)
            np.testing.assert_array_equal(a,models[f'k{k}_all_transitions'])
            np.testing.assert_array_equal(b,models[f'k{k}_stable_transitions'])
        result[unit]=dict(counts_verified=True,source_metadata_unchanged=True,medoids_from_real_learning_days=True,
                          all_100_omitted_date_predictions_replayed=True,transition_counts_verified=True,svg_count=15)
    core.save_json(out/'_data'/'verification.json',result)
    print('Seven-unit verification passed; 105 SVGs, three per unit/K',flush=True)


def bundle_code(out):
    files=[Path(__file__),Path(core.__file__),HERE/'test_seven_units.py',HERE/'jqrd_gt1'/'test_daily_modes.py']
    archive=out/'_data'/'reproducibility.zip'
    history={}
    original_hash=None
    if archive.exists():
        with zipfile.ZipFile(archive) as z:
            history={name:z.read(name) for name in z.namelist() if name.startswith('executed/')}
            entry=str(Path(__file__).relative_to(ROOT)).replace('\\','/')
            if entry in z.namelist():
                old=z.read(entry)
                original_hash=hashlib.sha256(old).hexdigest()
                if old!=Path(__file__).read_bytes():
                    history[f'executed/seven_units_{original_hash[:12]}.py']=old
    versions_path=out/'_data'/'code_versions.json'
    previous=core.read_json(versions_path) if versions_path.exists() else {}
    fitting_hash=previous.get('executed_fitting_entry_sha256',original_hash or core.digest(Path(__file__)))
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for name,content in history.items():
            z.writestr(name,content)
        for p in files:
            if p.exists():
                z.write(p,str(p.relative_to(ROOT)))
        manifest=core.read_json(out/'_data'/'manifest.json')
        z.writestr('requirements.txt','\n'.join(k+'=='+v for k,v in manifest['dependencies'].items())+'\n')
    versions={str(p.relative_to(ROOT)):core.digest(p) for p in files if p.exists()}
    versions['executed_fitting_entry_sha256']=fitting_hash
    core.save_json(versions_path,versions)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command',choices=['all','prepare','evaluate','report','verify'])
    parser.add_argument('--output',type=Path,default=DEFAULT_OUT)
    parser.add_argument('--workers',type=int,default=3)
    args=parser.parse_args()
    out=args.output.resolve()
    manifest=inventory_run(out)
    bundle_code(out)
    workers=max(1,min(args.workers,3))
    if args.command in ['all','prepare']:
        # Unit-level processes keep each CSV streaming and cap simultaneous memory usage.
        with ProcessPoolExecutor(max_workers=workers) as pool:
            jobs={pool.submit(prepare_unit,out,u,manifest['inventory'][u]):u for u in UNITS}
            for future in as_completed(jobs):
                future.result()
    if args.command in ['all','evaluate']:
        with ProcessPoolExecutor(max_workers=workers) as pool:
            jobs={pool.submit(evaluate_unit,out,u):u for u in UNITS}
            for future in as_completed(jobs):
                future.result()
    if args.command in ['all','report']:
        for unit in UNITS:
            report_unit(out,unit)
        summarize(out)
    if args.command in ['all','verify']:
        verify_outputs(out)
    print(f'{args.command} complete: {out}',flush=True)


if __name__=='__main__':
    main()
